from docx import Document

# Create a new document for the professional training delivery guide
doc = Document()
doc.add_heading('Professional Training Guide for Certificate-Level Networking Students', 0)

# Introduction
doc.add_heading('Introduction', level=1)
doc.add_paragraph(
    "This guide is designed to support hands-on practical training for certificate-level students "
    "with prior theoretical knowledge of computer networking. The focus is on real-world tasks, "
    "hardware handling, configuration, and troubleshooting techniques."
)

# Objectives
doc.add_heading('Training Objectives', level=1)
objectives = [
    "Identify and describe the functions of common networking components.",
    "Understand and handle different types of network cables and connectors.",
    "Identify LAN topologies and their applications.",
    "Assign and configure IP addresses on computers and networking devices.",
    "Terminate Ethernet cables using RJ-45 connectors and test them.",
    "Configure a basic router (e.g., TP-Link) for internet and LAN access.",
    "Troubleshoot network issues using structured diagnostic procedures."
]
for obj in objectives:
    doc.add_paragraph(obj, style='List Bullet')

# Training Modules
doc.add_heading('Training Modules & Activities', level=1)

# Module 1
doc.add_heading('Module 1: Identifying and Using Networking Components', level=2)
doc.add_paragraph(
    "• Components: Switch, Router, Patch Panel, Ethernet Cables (Cat5e/Cat6), RJ-45 connectors, Access Point, Network Tester.\n"
    "• Activity: Allow students to physically examine each item. Discuss use cases and demonstrate live setup."
)

# Module 2
doc.add_heading('Module 2: Network Cables and Connectors', level=2)
doc.add_paragraph(
    "• Focus: Understand straight-through vs crossover cables, fiber optic vs copper, RJ-45, RJ-11, SC, LC connectors.\n"
    "• Activity: Match cable types with use scenarios. Demonstrate cable crimping and verification."
)

# Module 3
doc.add_heading('Module 3: LAN Topologies', level=2)
doc.add_paragraph(
    "• Focus: Star, Bus, Ring, Mesh.\n"
    "• Activity: Use physical representations (wires, diagrams) to simulate each topology. Discuss pros and cons."
)

# Module 4
doc.add_heading('Module 4: Configuring IP Addresses', level=2)
doc.add_paragraph(
    "• Focus: Static and dynamic IP configuration.\n"
    "• Activity: Assign static IPs manually (Windows/Linux), and test communication using ping."
)

# Module 5
doc.add_heading('Module 5: Cable Termination', level=2)
doc.add_paragraph(
    "• Focus: Create Ethernet cables using crimping tools and RJ-45 connectors.\n"
    "• Activity: Students make their own cables, test with a cable tester, and connect PCs."
)

# Module 6
doc.add_heading('Module 6: Router Configuration (TP-Link)', level=2)
doc.add_paragraph(
    "• Focus: Set up internet via PPPoE/DHCP, Wi-Fi, and LAN DHCP.\n"
    "• Activity: Students reset, access the web GUI, configure SSID, IP range, and test connectivity."
)

# Module 7
doc.add_heading('Module 7: Network Troubleshooting', level=2)
doc.add_paragraph(
    "• Focus: Identify and resolve network problems.\n"
    "• Tools: ping, tracert, ipconfig/ifconfig, cable tester.\n"
    "• Activity: Diagnose common issues (IP conflict, disconnection, DNS failure)."
)

# Delivery Approach
doc.add_heading('Training Delivery Approach', level=1)
delivery_points = [
    "Each session begins with a 5-minute recap of related theory.",
    "Demonstration by instructor followed by guided student practice.",
    "Students work in pairs to promote collaboration and peer learning.",
    "Use of real hardware and network setups in each session.",
    "Assessment through hands-on tasks and scenario-based questions."
]
for point in delivery_points:
    doc.add_paragraph(point, style='List Bullet')

# Save the document
delivery_guide_path = "pdfs/Networking_Practical_Training_Guide.docx"
doc.save(delivery_guide_path)

delivery_guide_path
